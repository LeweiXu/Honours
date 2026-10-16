"""Refresh the empirical seminar slides in the editable final PPTX."""

from __future__ import annotations

import copy
import re
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
MPVRDU = Path('/home/lingwei/mpvrdu')
SOURCE = HERE / 'Seminar_final_before_10_15.pptx'
OUTPUT = HERE / 'Seminar_final.pptx'
ASSETS = HERE / 'figures/final'
BLUE = RGBColor(26, 86, 176)
NAVY = RGBColor(31, 78, 121)
BLACK = RGBColor(20, 24, 30)
MUTED = RGBColor(82, 91, 104)


def inch(value: float) -> int:
    return Inches(value)


def add_text(slide, text, x, y, w, h, size=14, bold=False, color=BLACK,
             align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split('\n')):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def clear(slide):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape._element)


def clone_shape(source_shape, dest_slide):
    element = copy.deepcopy(source_shape.element)
    dest_slide.shapes._spTree.insert_element_before(element, 'p:extLst')
    return dest_slide.shapes[-1]


def header(slide, template, title, section='Chapter 3 · Empirical study · Under review'):
    for i in (0, 1, 21):
        clone_shape(template.shapes[i], slide)
    title_shape = slide.shapes[-3]
    title_shape.text_frame.clear()
    title_shape.text_frame.paragraphs[0].text = title
    for run in title_shape.text_frame.paragraphs[0].runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(23)
        run.font.bold = True
        run.font.color.rgb = BLACK
    section_shape = slide.shapes[-2]
    section_shape.text_frame.clear()
    section_shape.text_frame.paragraphs[0].text = section
    for run in section_shape.text_frame.paragraphs[0].runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def add_image(slide, path, x, y, w, h=None):
    kwargs = {'width': inch(w)}
    if h is not None:
        kwargs['height'] = inch(h)
    return slide.shapes.add_picture(str(path), inch(x), inch(y), **kwargs)


def notes(slide, paragraphs):
    slide.notes_slide.notes_text_frame.text = '\n'.join(paragraphs)


def metric_from_row(line, index):
    cells = line.split('&')
    value = re.sub(r'\\cellcolor\{[^}]+\}', '', cells[index])
    value = re.sub(r'\\textbf\{([^}]+)\}', r'\1', value)
    return float(re.search(r'-?\d+(?:\.\d+)?', value).group())


def render_accuracy_table():
    source = (MPVRDU / 'tables/paper/abstention_rungs.tex').read_text()
    rows = source.splitlines()
    chosen = [
        ('MMLongBench-Doc', 'Chart', 1),
        ('MMLongBench-Doc', 'Figure', 1),
        ('LongDocURL', 'Text', 2),
        ('LongDocURL', 'Table', 2),
    ]
    extracted = []
    for corpus, label, section in chosen:
        marker = '\\textbf{' + corpus + '}'
        start = next(i for i, row in enumerate(rows) if marker in row)
        if section == 2:
            start = next(i for i, row in enumerate(rows) if marker in row)
        line = next(row for row in rows[start + 1:] if row.strip().startswith('\\quad ' + label + ' &'))
        tl, tlv = metric_from_row(line, 6), metric_from_row(line, 7)
        extracted.append((corpus, label, tl, tlv, tlv - tl))
    lines = [
        '% Values extracted from mpvrdu/tables/paper/abstention_rungs.tex.',
        '\\begin{table}[t]', '\\centering', '\\large',
        '\\begin{tabular}{@{}llrrr@{}}', '\\toprule',
        'Document & Evidence & Text & Text + image & Gain \\\\', '\\midrule',
    ]
    for corpus, label, tl, tlv, gain in extracted:
        short = 'MMLongBench' if corpus == 'MMLongBench-Doc' else corpus
        lines.append(f'{short} & {label} & {tl:.1f} & {tlv:.1f} & {gain:+.1f} \\\\')
    lines += ['\\bottomrule', '\\end{tabular}', '\\end{table}']
    tex = ASSETS / 'representation_accuracy.tex'
    tex.write_text('\n'.join(lines) + '\n')
    subprocess.run(['python3', str(HERE.parent / 'table_png.py'), str(tex),
                    '--out', str(ASSETS / 'representation_accuracy.png'),
                    '--dpi', '400', '--width', '16cm'], check=True, cwd=HERE.parent)


def prepare_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)
    subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-r', '175', '-png',
                    str(HERE.parent / 'Resources/empirical.pdf'),
                    str(ASSETS / 'empirical_page1')], check=True)
    subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-scale-to', '1500', '-png',
                    str(MPVRDU / 'figures/paper/distractor_exposure.pdf'),
                    str(ASSETS / 'distractor_exposure')], check=True)
    subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-scale-to', '1400', '-png',
                    str(MPVRDU / 'figures/paper/prompt_calibration.pdf'),
                    str(ASSETS / 'prompt_calibration')], check=True)
    subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-scale-to', '1200', '-png',
                    str(MPVRDU / 'figures/paper/fidelity.pdf'),
                    str(ASSETS / 'fidelity')], check=True)
    image = Image.open(ASSETS / 'fidelity.png')
    image.crop((0, round(image.height * 0.514), image.width, image.height)).save(
        ASSETS / 'fidelity_resolution_crop.png')
    render_accuracy_table()
    subprocess.run(['python3', str(HERE.parent / 'table_png.py'),
                    str(MPVRDU / 'tables/paper/encoding_cost.tex'),
                    '--out', str(ASSETS / 'encoding_cost_selected.png'),
                    '--dpi', '400', '--width', '16cm', '--rows', '1-2,5-6'],
                   check=True, cwd=HERE.parent)


def update_slide_10(slide, survey_slide):
    clear(slide)
    title = clone_shape(survey_slide.shapes[0], slide)
    title.text_frame.clear()
    run = title.text_frame.paragraphs[0].add_run()
    run.text = 'Locating Failure in Document Understanding'
    run.font.name = 'Calibri'
    run.font.size = Pt(21)
    run.font.bold = True
    run.font.color.rgb = BLACK
    chapter = clone_shape(survey_slide.shapes[1], slide)
    chapter.text_frame.clear()
    run = chapter.text_frame.paragraphs[0].add_run()
    run.text = 'Chapter 3 · Empirical study'
    run.font.name = 'Calibri'
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    add_image(slide, ASSETS / 'empirical_page1.png', 0.45, 1.27, 3.08, 3.71)
    clone_shape(survey_slide.shapes[3], slide)
    add_text(slide, 'An Empirical Attribution · Under review', 3.98, 1.45, 5.57, 0.36, 17, True, BLUE)
    add_text(slide, 'RESEARCH GAP', 3.98, 2.15, 5.57, 0.24, 9, True, NAVY)
    add_text(slide, 'A wrong answer does not tell us whether evidence was lost, missed, or misused.',
             3.98, 2.42, 5.57, 0.68, 15)
    add_text(slide, 'CONTRIBUTIONS', 3.98, 3.28, 5.57, 0.24, 9, True, NAVY)
    add_text(slide, 'Three failure loci: representation, selection, reasoning.\n'
             'Controlled interventions across two datasets and two model families.\n'
             'Deployment guidance from accuracy and cost.',
             3.98, 3.58, 5.57, 1.25, 13)
    clone_shape(survey_slide.shapes[9], slide)
    notes(slide, [
        'The survey maps how systems manage evidence. This second paper asks where those systems fail.',
        'A final wrong answer hides three different causes. The relevant information may be lost while encoding a page, a needed page may be omitted during selection, or the model may fail to use evidence it has received.',
        'We separate these into representation, selection and reasoning. We test each through controlled interventions, then check whether the findings transfer across MMLongBench-Doc, LongDocURL, Qwen3-VL and Gemma-3.',
    ])


def update_slide_11(slide, framework_image, template):
    clear(slide)
    header(slide, template, 'Attribution by Construction')
    add_text(slide, 'Change one stage while holding the other two fixed.',
             0.45, 0.78, 9.1, 0.30, 15, True, BLUE)
    add_image(slide, framework_image, 0.55, 1.20, 8.90, 2.68)
    cards = [
        (0.55, 'REPRESENTATION', 'T · TL · TLV · V\nParsers and image resolution'),
        (3.58, 'SELECTION', 'Gold pages as control\nDrop evidence or add distractors'),
        (6.61, 'REASONING', 'Qwen3-VL and Gemma-3\nModel scale and prompt modes'),
    ]
    for x, label, body in cards:
        add_text(slide, label, x, 4.05, 2.85, 0.24, 10, True, NAVY)
        add_text(slide, body, x, 4.32, 2.85, 0.56, 12)
    add_text(slide, 'MMLongBench-Doc + LongDocURL  ·  Accuracy judge for correctness; rubric judge for refusal',
             0.55, 5.00, 8.9, 0.22, 9, False, MUTED)
    notes(slide, [
        'The framework is also our experimental setup. Every document question-answering system has to represent pages, choose what reaches the model and reason over that evidence.',
        'We use one encode, retrieve and answer pipeline. For representation we vary embedded text, parser text and page images, including parser and resolution choices. For selection we start with annotated gold pages, then omit required pages or add non-gold pages. For reasoning we vary the backbone and its instructions.',
        'Each comparison changes one stage while the others stay fixed. MMLongBench-Doc supplies evidence-page and evidence-source annotations; LongDocURL checks whether the patterns transfer to born-digital documents. Correctness and refusal come from their respective judges.',
    ])


def update_slide_12(slide, template):
    clear(slide)
    header(slide, template, 'Finding 1 · Representation')
    add_text(slide, 'Use vision where the task needs it; text helps low-resolution vision recover detail.',
             0.45, 0.77, 9.1, 0.34, 14, True, BLUE)
    add_text(slide, 'Accuracy by evidence type (%)', 0.45, 1.18, 5.0, 0.23, 10, True, NAVY)
    add_image(slide, ASSETS / 'representation_accuracy.png', 0.45, 1.45, 5.03)
    add_text(slide, 'Cost of vision on two V100s', 0.45, 3.02, 5.0, 0.23, 10, True, NAVY)
    add_image(slide, ASSETS / 'encoding_cost_selected.png', 0.45, 3.29, 4.80)
    add_text(slide, 'Low-resolution image + parser text', 5.75, 1.18, 3.80, 0.23, 10, True, NAVY)
    add_image(slide, ASSETS / 'fidelity_resolution_crop.png', 5.56, 1.49, 4.00)
    add_text(slide, '49.2% at low resolution vs 54.7% at high resolution',
             5.62, 3.65, 3.93, 0.48, 13, True, BLUE)
    add_text(slide, 'That keeps 90% of the high-resolution accuracy on the same 847 questions.',
             5.62, 4.17, 3.93, 0.54, 11)
    add_text(slide, 'T embedded text · TL parser text · TLV parser text + image · V image only',
             0.45, 5.08, 9.10, 0.18, 9, False, MUTED)
    notes(slide, [
        'Vision is valuable when the answer depends on appearance. On MMLongBench-Doc chart questions, parser text reaches 19.1 percent accuracy and text with the page image reaches 45.5 percent. Figure questions show a similar gain. On born-digital LongDocURL, adding the image to parser text gives only 2.8 points on text questions and is slightly worse on table questions.',
        'Vision also costs time. On two 16 GB V100s, an image-only input takes 22.7 seconds to prefill against 1.1 seconds for embedded text, about 21 times as long, even though both have about 900 input tokens. Text with images takes 24.2 seconds and has 11.2 percent out-of-memory failures, against 1.2 percent for text.',
        'We can often render the image cheaply. Parser text plus a low-resolution image reaches 49.2 percent accuracy, compared with 54.7 percent at high resolution, on the same 847 questions. That keeps 90 percent of the high-resolution result. Text compensates especially on text and table questions; visual detail remains more important for charts and figures.',
    ])


def update_slide_13(slide, template):
    clear(slide)
    header(slide, template, 'Finding 2 · Selection')
    add_text(slide, 'All required pages stay in context, yet extra pages change the answer.',
             0.45, 0.77, 9.1, 0.34, 14, True, BLUE)
    add_image(slide, ASSETS / 'distractor_exposure.png', 0.45, 1.35, 9.10)
    add_text(slide, '+3 pages: Qwen3-VL multi-hop accuracy 43.9% → 35.1%',
             0.50, 4.95, 5.55, 0.28, 13, True, BLUE)
    add_text(slide, '13.3% broken · 7.0% repaired (+3)',
             6.00, 4.95, 3.55, 0.28, 11, True, NAVY)
    notes(slide, [
        'The distractor experiment keeps every annotated gold page in the context. We then add the highest-ranked non-gold pages, so evidence coverage stays complete.',
        'The top row tracks accuracy. With Qwen3-VL-8B, multi-hop accuracy falls from 43.9 to 35.1 percent after just three extra pages, a drop of 8.8 points. Gemma-3 shows the same early decline. The bottom row tracks paired answer flips: at plus three, 13.3 percent of Qwen3-VL answers break while 7.0 percent are repaired.',
        'The repairs matter. Added context can provide useful clues even when it is not annotated as gold. But the larger broken share means the net effect is harmful, particularly for multi-hop questions. Selection must control distractor exposure as well as evidence coverage.',
    ])


def update_slide_14(slide, template):
    clear(slide)
    header(slide, template, 'Finding 3 · Reasoning')
    add_text(slide, 'Prompts move the accuracy and refusal trade-off.',
             0.45, 0.77, 9.1, 0.34, 14, True, BLUE)
    add_image(slide, ASSETS / 'prompt_calibration.png', 0.45, 1.27, 5.40)
    add_text(slide, 'Qwen3-VL-8B', 6.10, 1.48, 3.4, 0.30, 14, True, NAVY)
    add_text(slide, 'Explicit abstention raises correct refusal on unanswerable questions:',
             6.10, 1.92, 3.35, 0.75, 13)
    add_text(slide, '45.5% → 74.6%', 6.10, 2.72, 3.35, 0.34, 18, True, BLUE)
    add_text(slide, 'It also raises false refusal on answerable questions:',
             6.10, 3.27, 3.35, 0.75, 13)
    add_text(slide, '6.2% → 27.7%', 6.10, 4.05, 3.35, 0.34, 18, True, BLUE)
    add_text(slide, 'Reasoning before abstaining recovers answerable accuracy to 53.0%.',
             6.10, 4.55, 3.35, 0.55, 11)
    notes(slide, [
        'The third failure point is reasoning over evidence the model already has. A fluent answer can rest on a premise the document does not support.',
        'This figure compares four prompt modes. For Qwen3-VL-8B, telling the model it may abstain raises correct refusal on unanswerable questions from 45.5 to 74.6 percent. It also increases false refusal on answerable questions from 6.2 to 27.7 percent, and answerable accuracy drops from 52.5 to 41.4 percent.',
        'Adding a reasoning step before the abstention decision brings answerable accuracy back to 53.0 percent and raises correct unanswerable refusal to 85.5 percent for Qwen3-VL. Gemma-3 does not get the same benefit. The final strategy comparison needs rubric-judged refusal rates and a clear account of its unmatched precision and prompt settings.',
    ])


def update_slide_15_notes(slide):
    notes(slide, [
        'The paper also checks individual questions to show what each failure looks like. Representation can hide an icon colour or corrupt a table during parsing. Selection can omit a required page or bury it among competing pages. Reasoning can accept a false premise or refuse despite sufficient evidence.',
        'The design responses follow the same three stages: add the image when appearance matters or text is unreliable; retrieve broadly, then filter the evidence as a set; and calibrate the reasoner against the evidence it actually received.',
    ])


def main():
    if not SOURCE.exists():
        raise SystemExit(f'Missing preserved source deck: {SOURCE}')
    prepare_assets()
    prs = Presentation(SOURCE)
    if len(prs.slides) != 22:
        raise SystemExit(f'Expected 22 source slides, found {len(prs.slides)}')
    survey = prs.slides[4]
    template = prs.slides[11]
    framework = next(shape for shape in prs.slides[10].shapes if shape.shape_type == 13)
    framework_path = ASSETS / 'attribution_framework.png'
    framework_path.write_bytes(framework.image.blob)
    update_slide_10(prs.slides[9], survey)
    update_slide_11(prs.slides[10], framework_path, template)
    update_slide_12(prs.slides[12], template)
    update_slide_13(prs.slides[13], template)
    update_slide_14(prs.slides[14], template)
    update_slide_15_notes(prs.slides[15])
    slide_id = prs.slides._sldIdLst[11]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)
    prs.save(OUTPUT)
    print(f'Updated {OUTPUT}: {len(prs.slides)} slides')


if __name__ == '__main__':
    main()
