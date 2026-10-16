"""Place the full representation table and specified paper figures in the findings slides."""

from pathlib import Path

from pptx import Presentation

from revise_final_slides_v2 import ASSETS, BLUE, clear, header, notes, picture, textbox
from revise_final_revised_v3 import bullet

HERE = Path(__file__).resolve().parent
REPO = Path('/home/lingwei/mpvrdu')
SOURCE = HERE / 'Seminar_final_revised.pptx'
OUTPUT = HERE / 'Seminar_final_revised_checked.pptx'
SECTION = 'Chapter 3 · Empirical study · Under review'


def representation(slide, template):
    clear(slide)
    header(slide, template, 'Finding 1 · Representation', SECTION)
    textbox(slide, 'Vision is task-dependent; text helps recover detail at low image resolution.',
            .45, .77, 9.1, .34, 14, BLUE, True)
    picture(slide, ASSETS / 'abstention_rungs_full.png', .55, 1.22, 4.05)
    picture(slide, ASSETS / 'fidelity_paper.png', 5.25, 1.22, 3.90)
    bullet(slide, 'Vision adds 26.4 points on charts, but little on born-digital text and tables; image prefill costs about 21× text.',
           .50, 4.73, 9.0, .22, 9.8)
    bullet(slide, 'Text + low-res images: 49.2% accuracy, retaining 90% of the high-res result (54.7%).',
           .50, 4.99, 9.0, .22, 10.3, BLUE, True)
    notes(slide, [
        'The left visual is the full Table 1, including both datasets, every evidence-source row and all four representations. T is embedded text, TL parser text, TLV parser text with the page image, and V image only. Accuracy and abstention are separate judged outcomes.',
        'The modality ceiling depends on the task and the document. For MMLongBench-Doc charts, parser text scores 19.1 percent and text with the image reaches 45.5 percent. Figures rise from 23.8 to 44.8 percent. Born-digital LongDocURL text gains only 2.8 points, and tables move from 69.5 to 68.2 percent.',
        'The right visual is the complete conversion-fidelity figure from the paper. Its first panel shows how page images compensate for different parsers on digital and scanned documents. The second panel shows how text compensates for low-resolution images, especially for text and table questions.',
        'At low resolution, text plus images reaches 49.2 percent overall on the same 847 answerable questions, 90 percent of the high-resolution result of 54.7 percent. Vision has a cost: image-only prefill takes 22.7 seconds against 1.1 seconds for embedded text on two 16 GB V100s, roughly 21 times as long.',
    ])


def selection(slide, template):
    clear(slide)
    header(slide, template, 'Finding 2 · Selection', SECTION)
    textbox(slide, 'Complete evidence coverage does not protect an answer from distractors.',
            .45, .77, 9.1, .34, 14, BLUE, True)
    picture(slide, REPO / 'figures/paper/png/distractor_exposure_merged.png', .50, 1.22, 4.30)
    textbox(slide, 'ADD k NON-GOLD PAGES', 5.23, 1.40, 4.2, .25, 10, BLUE, True)
    bullet(slide, 'All gold pages remain. The added pages are the highest-ranked non-gold pages.',
           5.23, 1.77, 4.2, .65, 13)
    bullet(slide, '+3 pages cuts multi-hop accuracy: Qwen3 43.9 → 35.1%; Gemma-3 35.4 → 28.7%.',
           5.23, 2.61, 4.2, .82, 13, BLUE, True)
    bullet(slide, 'At +3, 13.3% of Qwen3 answers break and 7.0% repair. Further additions keep flipping answers.',
           5.23, 3.62, 4.2, .91, 13)
    textbox(slide, 'Selection must control distractor exposure as well as retain evidence.',
            5.23, 4.64, 4.2, .42, 11, BLUE, True)
    notes(slide, [
        'This is distractor_exposure_merged.pdf, the unified MMLongBench-Doc plot. Colour separates single-hop from multi-hop questions. Solid lines and filled markers are Qwen3-VL-8B; dashed lines and open markers are Gemma-3-12B. The lower panel uses filled and outline bars for the two models.',
        'Every gold page remains in context. We add the top k retrieved non-gold pages, so evidence recall is fixed. At plus three, multi-hop accuracy falls from 43.9 to 35.1 percent for Qwen3 and from 35.4 to 28.7 percent for Gemma-3.',
        'The lower panel reports paired verdict flips at each addition step. At plus three, 13.3 percent of Qwen3 answers break and 7.0 percent repair. Non-gold context can supply useful clues, but more answers are derailed than repaired. Later steps keep producing flips, while multi-hop accuracy remains below the gold-only baseline. This makes distractor exposure a separate selection problem from missing evidence.',
    ])


def reasoning(slide, template):
    clear(slide)
    header(slide, template, 'Finding 3 · Reasoning', SECTION)
    textbox(slide, 'Prompting moves the refusal trade-off; additional inference gives selective gains.',
            .45, .77, 9.1, .34, 13.5, BLUE, True)
    picture(slide, REPO / 'figures/paper/png/prompt_calibration.png', .45, 1.26, 4.45)
    picture(slide, REPO / 'figures/paper/png/reasoning_strategies.png', 5.30, 1.26, 4.25)
    bullet(slide, 'Abstain raises correct refusal, but costs 11.1 points of answerable accuracy for Qwen3.',
           .45, 4.66, 4.45, .48, 10.5)
    bullet(slide, 'Answerable accuracy: 60.6% with SC and 55.1% with Reflexion, vs 52.5% plain.',
           5.30, 4.02, 4.25, .53, 11, BLUE, True)
    bullet(slide, 'Reflexion correctly refuses 61.5%, vs 45.5% plain.',
           5.30, 4.65, 4.25, .47, 11)
    notes(slide, [
        'The left figure is the original prompt_calibration.png: four prompt modes for Qwen3 and Gemma-3. The separate right figure compares only self-consistency and Reflexion on Qwen3-VL-8B at 4-bit.',
        'Both figures use the same outcome definitions. Grey answerable segments are correctness from the accuracy judge. Blue answerable segments are refusal from the taxonomy judge. On unanswerable questions, blue is the refusal credited by the taxonomy judge.',
        'For Qwen3, the abstain instruction raises correct unanswerable refusal from 45.5 to 74.6 percent and answerable refusal from 6.2 to 27.7 percent, but reduces answerable accuracy from 52.5 to 41.4 percent. Combining CoT with abstention recovers answerable accuracy to 53.0 percent and reaches 85.5 percent correct refusal. Gemma does not show that recovery.',
        'Self-consistency reaches 60.6 percent answerable accuracy, about eight points above plain prompting; Reflexion reaches 55.1 percent, about 2.6 points above it. Their correct unanswerable refusal rates are 53.3 and 61.5 percent. Reflexion is 16 points above the plain 45.5 percent baseline on this taxonomy-judged measure.',
        'These are observed comparisons: Qwen3 strategy runs use 4-bit precision and different prompts, while the original Qwen3 prompt modes use bf16. Gemma prompt modes use 4-bit. Reflexion answerable refusal is provisional at 294 rubric-judged cases; self-consistency has 844. The figure prints these counts.',
    ])


if __name__ == '__main__':
    deck = Presentation(SOURCE)
    template = Presentation(SOURCE).slides[11]
    assert len(deck.slides) == 21
    representation(deck.slides[11], template)
    selection(deck.slides[12], template)
    reasoning(deck.slides[13], template)
    deck.save(OUTPUT)
    print(OUTPUT)
