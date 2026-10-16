"""Put notes.md into the presenter notes of Seminar_final.pptx.

    ~/miniconda3/bin/python update_notes.py            # needs python-pptx

The text under each `## Slide N: ...` heading becomes the notes of slide N, one paragraph
per line, with the bullet dashes and bold markers taken off. A slide with no heading in
notes.md keeps the notes it has. Close the deck in PowerPoint first.
"""
import re, sys
from pathlib import Path
from pptx import Presentation

SEM = Path(__file__).parent
PPTX = SEM / "Seminar_final.pptx"
NOTES = SEM / "notes.md"

def read_notes(path):
    """Slide number -> its lines, the same reading build.py gives notes.md."""
    out = {}
    for sec in re.split(r"(?m)^(?=#{1,2} )", path.read_text(encoding="utf-8")):
        m = re.match(r"## Slide (\d+):", sec)
        if not m:
            continue
        lines = []
        for ln in sec.splitlines()[1:]:
            ln = re.sub(r"<!--.*?-->", "", ln).strip()
            if not ln or ln == "---":
                continue
            ln = re.sub(r"^- ", "", ln)
            ln = re.sub(r"\*\*(.+?)\*\*", r"\1", ln)
            lines.append(ln)
        out[int(m[1])] = lines
    return out

def main():
    notes = read_notes(NOTES)
    prs = Presentation(PPTX)
    n = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        if i in notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(notes[i])
    extra = sorted(k for k in notes if k > n)
    missing = [i for i in range(1, n + 1) if i not in notes]
    prs.save(PPTX)
    print(f"{PPTX.name}: notes written to {n - len(missing)} of {n} slides")
    if missing:
        print("no section in notes.md, left as they were: " + ", ".join(map(str, missing)))
    if extra:
        print("in notes.md but past the last slide, ignored: " + ", ".join(map(str, extra)))

if __name__ == "__main__":
    sys.exit(main())
