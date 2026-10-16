"""Update the seminar findings while keeping the deck's native layouts."""

from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SOURCE = HERE / 'Seminar_final.pptx'
OUTPUT = HERE / 'Seminar_final_revised.pptx'
ASSETS = HERE / 'figures/final'
INK = RGBColor(28, 36, 49)
BLUE = RGBColor(24, 79, 169)
MUTED = RGBColor(88, 99, 112)


def clone(shape, slide):
    elem = copy.deepcopy(shape.element)
    slide.shapes._spTree.insert_element_before(elem, 'p:extLst')
    return slide.shapes[-1]


def clear(slide):
    for shape in list(slide.shapes):
        slide.shapes._spTree.remove(shape.element)


def replace_runs(shape, parts):
    para = shape.text_frame.paragraphs[0]
    while len(para.runs) < len(parts):
        para.add_run()
    for run, value in zip(para.runs, parts):
        run.text = value
    for run in para.runs[len(parts):]:
        run.text = ''


def textbox(slide, text, x, y, w, h, size=12, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split('\n')):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.name = 'Calibri'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def picture(slide, path, x, y, w):
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))


def notes(slide, lines):
    slide.notes_slide.notes_text_frame.text = '\n'.join(lines)


def header(slide, exemplar, title, subtitle):
    for ix in (0, 1, 2):
        clone(exemplar.shapes[ix], slide)
    replace_runs(slide.shapes[0], [title])
    replace_runs(slide.shapes[1], [subtitle])


def slide10(slide, survey):
    clear(slide)
    title = clone(survey.shapes[0], slide)
    replace_runs(title, ['Locating Failure in Multi-Page ',
                         'Document Understanding'])
    chapter = clone(survey.shapes[1], slide)
    replace_runs(chapter, ['Chapter 3 · Empirical study', ''])
    page = picture(slide, ASSETS / 'empirical_page1_crop.png', 0.45, 1.27, 3.08)
    slide.shapes._spTree.remove(page.element)
    slide.shapes._spTree.insert(4, page.element)
    clone(survey.shapes[3], slide)
    textbox(slide, 'UNDER REVIEW', 7.42, 0.83, 2.13, 0.32, 16, BLUE, True)
    for ix in (5, 6, 7, 8, 9, 10):
        clone(survey.shapes[ix], slide)
    replace_runs(slide.shapes[5], ['RESEARCH GAP'])
    replace_runs(slide.shapes[6], ['A final wrong answer hides ',
                                    'where evidence fails',
                                    ': encoding, selection, or reasoning.'])
    replace_runs(slide.shapes[7], ['CONTRIBUTIONS'])
    body = slide.shapes[8].text_frame.paragraphs
    for run in body[0].runs: run.text = ''
    body[0].runs[0].text = 'Three architecture-independent '
    body[0].runs[1].text = 'failure loci'
    for run in body[1].runs: run.text = ''
    body[1].runs[0].text = 'Controlled interventions across two benchmarks and two model families'
    for run in body[2].runs: run.text = ''
    body[2].runs[0].text = 'Deployment guidance from accuracy and cost'
    replace_runs(slide.shapes[10], ['First-authored empirical study, ',
                                     'under review', '', ''])
    notes(slide, [
        'The survey in Chapter 2 maps how systems manage evidence. This empirical paper asks where the evidence pipeline fails.',
        'A final wrong answer does not reveal whether a page was encoded poorly, the needed evidence was not selected, or the model failed to reason over evidence it received.',
        'We isolate representation, selection and reasoning with controlled interventions on MMLongBench-Doc and LongDocURL, using Qwen3-VL and Gemma-3 backbones. The study is under review.',
    ])


def slide12(slide, exemplar):
    clear(slide)
    header(slide, exemplar, 'Finding 1 · Representation',
           'Chapter 3 · Empirical study · Under review')
    textbox(slide, 'Vision is task-dependent; text rescues low-resolution images.',
            .45, .77, 9.1, .34, 14, BLUE, True)
    picture(slide, ASSETS / 'representation_table.png', .45, 1.33, 4.69)
    picture(slide, ASSETS / 'fidelity_paper.png', 5.47, 1.28, 3.62)
    textbox(slide, 'Chart: 19.1% text vs 45.5% text + image. LongDocURL tables gain no accuracy.',
            .45, 4.20, 4.70, .47, 11)
    textbox(slide, 'Image-only prefill: 22.7 s vs 1.1 s for embedded text, about 21× longer.',
            .45, 4.68, 4.70, .34, 11, BLUE, True)
    textbox(slide, 'Text + low-res image reaches 49.2%, or 90% of the high-res result (54.7%).',
            5.28, 4.67, 4.27, .37, 10.5, BLUE, True)
    textbox(slide, 'T embedded text · TL parser text · TLV text + image · V image only',
            .45, 5.08, 9.0, .15, 8.5, MUTED)
    notes(slide, [
        'The LaTeX table and conversion-fidelity figure are rendered from the paper sources. The table shows that representation depends on both the task and the document type.',
        'For MMLongBench-Doc chart questions, parser text scores 19.1 percent and parser text plus the page image scores 45.5 percent. Figure questions gain 21 points. In born-digital LongDocURL, text questions gain only 2.8 points with the image, and table questions fall from 69.5 to 68.2 percent.',
        'Vision costs substantially more. On two 16 GB V100s, image-only input takes 22.7 seconds to prefill, compared with 1.1 seconds for embedded text, about 21 times as long. Text plus image takes 24.2 seconds and has 11.2 percent out-of-memory failures. These are measured input modes, not equal-accuracy configurations.',
        'The paper fidelity figure shows parser spread and resolution effects. Text plus low-resolution page images scores 49.2 percent overall on 847 answerable questions, 90 percent of the 54.7 percent high-resolution result. Text rescues detail, especially in text and table questions; appearance still matters for charts and figures.',
    ])


def slide13(slide, exemplar):
    clear(slide)
    header(slide, exemplar, 'Finding 2 · Selection',
           'Chapter 3 · Empirical study · Under review')
    textbox(slide, 'Distractors change answers even when every gold page remains in context.',
            .45, .77, 9.1, .34, 14, BLUE, True)
    picture(slide, ASSETS / 'distractor_paper.png', .45, 1.32, 9.10)
    textbox(slide, '+3 pages: multi-hop accuracy falls 43.9 → 35.1% (Qwen3) and 35.4 → 28.7% (Gemma-3).',
            .45, 4.76, 9.1, .24, 11, BLUE, True)
    textbox(slide, 'Qwen3 flips at +3: 13.3% break, 7.0% repair; further distractors keep changing answers.',
            .45, 5.03, 9.1, .20, 9.5, MUTED)
    notes(slide, [
        'We keep every annotated gold page and add the highest-ranked non-gold pages. Evidence coverage remains complete, so the change isolates distractor exposure.',
        'With three extra pages, Qwen3-VL multi-hop accuracy falls from 43.9 to 35.1 percent. Gemma-3 falls from 35.4 to 28.7 percent. The full paper figure shows both model families and LongDocURL.',
        'The paired flips are revealing: at plus three, 13.3 percent of Qwen3 answers break and 7.0 percent repair. Additional pages can sometimes supply useful clues, but they more often derail an answer. The effect continues at later k values, so selection has to manage extra context as well as retain gold evidence.',
    ])


def slide14(slide, exemplar):
    clear(slide)
    header(slide, exemplar, 'Finding 3 · Reasoning',
           'Chapter 3 · Empirical study · Under review')
    textbox(slide, 'Prompting trades accuracy for refusal; inference strategies shift the balance.',
            .45, .77, 9.1, .34, 14, BLUE, True)
    picture(slide, ASSETS / 'reasoning_comparison.png', .45, 1.15, 9.10)
    textbox(slide, 'Abstain prompt: unanswerable correct 47.5 → 74.6%, answerable correct 52.5 → 41.4%.',
            .45, 4.53, 9.1, .23, 10.5, BLUE, True)
    textbox(slide, 'Self-consistency: +8.0 pt answerable vs plain. Reflexion: +2.5 pt answerable, +13.9 pt unanswerable.',
            .45, 4.80, 9.1, .23, 10.3)
    textbox(slide, 'Reflexion abstention uses 294 judged answerable cases; strategy runs use 4-bit precision.',
            .45, 5.07, 9.1, .15, 8.3, MUTED)
    notes(slide, [
        'This figure uses current accuracy-judge results for correct answers and the separate rubric judge for answerable abstentions. Each bar is a percentage of successful judged cases. Plain, CoT and abstain have 847 answerable accuracy cases and 244 unanswerable cases; the rubric judge has 833 answerable cases for those modes.',
        'The abstain instruction lifts unanswerable correctness from 47.5 to 74.6 percent, but reduces answerable correctness from 52.5 to 41.4 percent and raises answerable abstention from 6.5 to 27.7 percent. CoT plus abstain recovers answerable correctness to 53.0 percent and reaches 85.5 percent on unanswerable questions, among 779 and 234 completed cases respectively.',
        'Self-consistency reaches 60.6 percent answerable correctness, eight points over plain prompting, but 54.5 percent unanswerable correctness. Reflexion reaches 55.1 and 61.5 percent, up 2.5 and 13.9 points over plain. These are observed differences, not a matched causal estimate: the two strategy runs use 4-bit precision and different prompting, whereas the four prompt modes use bf16.',
        'Reflexion answerable abstention is provisional: only 294 of its 846 answerable accuracy cases have rubric judgements so far. The figure labels this limitation on the slide. Self-consistency has 844 rubric-judged answerable cases.',
    ])


def slide17(slide, exemplar):
    clear(slide)
    header(slide, exemplar, 'Simple Method: Results',
           'Chapter 4 · Method & conclusion')
    textbox(slide, 'Qwen3-VL-8B comparisons across three document benchmarks.',
            .45, .77, 9.1, .34, 14, BLUE, True)
    picture(slide, ASSETS / 'method_qwen3_docbench.png', .45, 1.30, 9.10)
    textbox(slide, 'Ours: 56.5% MMLongBench-Doc · 61.3% LongDocURL · 71.6% DocBench',
            .45, 4.34, 9.1, .31, 14, BLUE, True)
    textbox(slide, 'DocBench is a provisional 762/1,065 result; MARDoc reports 72.3%. Multi-page: 39.2 vs DocTrace 41.1%.',
            .45, 4.71, 9.1, .34, 10.5)
    notes(slide, [
        'This is the Qwen3-VL-8B section of the updated LaTeX comparison table, including DocBench. Ours uses one model at 4-bit without task-specific training.',
        'The method scores 56.5 percent on MMLongBench-Doc and 61.3 percent on LongDocURL under their official protocols. In DocBench the current 762 of 1,065 completed questions score 71.6 percent. MARDoc reports 72.3 percent there, so do not describe ours as best in that column.',
        'Our multi-page score is 39.2 percent, below DocTrace at 41.1 percent. DocBench evaluations in the source table also use a different evaluator for ours and the MARDoc results, so that column is a reported comparison with a protocol caveat.',
    ])


def main():
    prs = Presentation(SOURCE)
    assert len(prs.slides) == 21
    template = Presentation(SOURCE)
    survey, exemplar = template.slides[4], template.slides[11]
    slide10(prs.slides[9], survey)
    slide12(prs.slides[11], exemplar)
    slide13(prs.slides[12], exemplar)
    slide14(prs.slides[13], exemplar)
    slide17(prs.slides[16], exemplar)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
