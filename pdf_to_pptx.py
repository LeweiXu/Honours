#!/usr/bin/env python3
"""
pdf_to_pptx.py  --  Turn a slide PDF into a PowerPoint, one full-page image per slide.
No fonts or colors for PowerPoint to mis-render: each slide IS the rendered image.

Usage:  python pdf_to_pptx.py deck.pdf deck.pptx
"""

import sys, os, subprocess, tempfile, glob
from pptx import Presentation
from pptx.util import Inches

DPI = 200  # 200 is crisp for 16:9; bump to 300 if you want extra sharpness

def render_with_pymupdf(pdf_path, out_dir):
    """Preferred: PyMuPDF (pip install pymupdf). No external poppler needed."""
    import fitz  # PyMuPDF
    doc = fitz.open(pdf_path)
    paths = []
    zoom = DPI / 72.0
    mat = fitz.Matrix(zoom, zoom)
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=mat, alpha=False)
        p = os.path.join(out_dir, f"page_{i:03d}.png")
        pix.save(p)
        paths.append(p)
    return paths

def render_with_poppler(pdf_path, out_dir):
    """Fallback: pdftoppm (poppler-utils). 'brew install poppler' / apt install poppler-utils."""
    prefix = os.path.join(out_dir, "page")
    subprocess.run(["pdftoppm", "-png", "-r", str(DPI), pdf_path, prefix], check=True)
    return sorted(glob.glob(prefix + "*.png"))

def build_pptx(image_paths, out_pptx):
    prs = Presentation()
    # 16:9 widescreen canvas
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]  # fully blank layout
    for img in image_paths:
        slide = prs.slides.add_slide(blank)
        # full-bleed: fill the whole slide with the page image
        slide.shapes.add_picture(img, 0, 0,
                                 width=prs.slide_width,
                                 height=prs.slide_height)
    prs.save(out_pptx)

def main():
    if len(sys.argv) != 3:
        print("Usage: python pdf_to_pptx.py <input.pdf> <output.pptx>")
        sys.exit(1)
    pdf_path, out_pptx = sys.argv[1], sys.argv[2]

    with tempfile.TemporaryDirectory() as tmp:
        try:
            imgs = render_with_pymupdf(pdf_path, tmp)
            print(f"Rendered {len(imgs)} pages with PyMuPDF.")
        except ImportError:
            imgs = render_with_poppler(pdf_path, tmp)
            print(f"Rendered {len(imgs)} pages with pdftoppm.")
        if not imgs:
            print("No pages rendered — check the PDF path.")
            sys.exit(1)
        build_pptx(imgs, out_pptx)
    print(f"Wrote {out_pptx}  ({len(imgs)} slides)")

if __name__ == "__main__":
    main()
